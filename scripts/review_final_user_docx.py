from __future__ import annotations

import argparse
import difflib
import json
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def body_lines(doc: Document) -> list[str]:
    lines: list[str] = []
    paragraph_index = 0
    table_index = 0
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = doc.paragraphs[paragraph_index]
            text = " ".join(paragraph.text.split())
            if text:
                lines.append(f"P{paragraph_index:03d}|{text}")
            paragraph_index += 1
        elif child.tag == qn("w:tbl"):
            table = doc.tables[table_index]
            lines.append(f"TABLE{table_index}")
            for row_index, row in enumerate(table.rows):
                cells = [" ".join(cell.text.split()) for cell in row.cells]
                lines.append(f"T{table_index}R{row_index}|" + " || ".join(cells))
            table_index += 1
    return lines


def package_flags(path: Path) -> dict[str, object]:
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        document_xml = archive.read("word/document.xml")
        settings_xml = archive.read("word/settings.xml") if "word/settings.xml" in names else b""
        return {
            "zip_ok": archive.testzip() is None,
            "comments_part": "word/comments.xml" in names,
            "tracked_insertions": document_xml.count(b"<w:ins"),
            "tracked_deletions": document_xml.count(b"<w:del"),
            "track_revisions_setting": b"<w:trackRevisions" in settings_xml,
            "embedded_media": sorted(name for name in names if name.startswith("word/media/")),
        }


def section_summary(doc: Document) -> list[dict[str, float | int]]:
    result = []
    for section in doc.sections:
        result.append(
            {
                "width_in": round(section.page_width.inches, 3),
                "height_in": round(section.page_height.inches, 3),
                "top_in": round(section.top_margin.inches, 3),
                "bottom_in": round(section.bottom_margin.inches, 3),
                "left_in": round(section.left_margin.inches, 3),
                "right_in": round(section.right_margin.inches, 3),
            }
        )
    return result


def all_text(doc: Document) -> str:
    pieces = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            pieces.extend(cell.text for cell in row.cells)
    return "\n".join(pieces)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("current", type=Path)
    parser.add_argument("baseline", type=Path)
    args = parser.parse_args()

    current = Document(args.current)
    baseline = Document(args.baseline)
    current_text = all_text(current)
    diff = list(
        difflib.unified_diff(
            body_lines(baseline),
            body_lines(current),
            fromfile="baseline",
            tofile="current",
            lineterm="",
        )
    )

    checks = {
        "has_latest_scenario_label": "최신 승인 대표 시나리오(합성 기사 활용)" in current_text,
        "a100_300_hours": "활용 시간: 300시간" in current_text,
        "demo_possible_checked": "■ 시연 가능" in current_text,
        "demo_conditional_checked": "■ 조건부 가능" in current_text,
        "video_url_blank": bool(re.search(r"영상 URL:\s*(?:\n|$)", current_text)),
        "forbidden_photo_disclaimer_absent": "합성 시연, 실제 사고확률 아님" not in current_text,
        "effect_claim_boundary_absent": "효과 주장 경계" not in current_text,
        "no_public_demo_url": not bool(re.search(r"https?://", current_text)),
        "exact_eval_boundary": (
            "모든 평가 내 숫자는 합성된 Live 자격평가입니다." in current_text
            and "실제 사고 및 현장 생산성 향상으로 해석하지 않습니다." in current_text
        ),
        "no_authoring_slashes": "// 해당" not in current_text,
    }

    report = {
        "current": str(args.current),
        "baseline": str(args.baseline),
        "structure": {
            "paragraphs": len(current.paragraphs),
            "tables": len(current.tables),
            "inline_shapes": len(current.inline_shapes),
            "sections": section_summary(current),
        },
        "baseline_structure": {
            "paragraphs": len(baseline.paragraphs),
            "tables": len(baseline.tables),
            "inline_shapes": len(baseline.inline_shapes),
            "sections": section_summary(baseline),
        },
        "package": package_flags(args.current),
        "checks": checks,
        "diff": diff,
    }
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
