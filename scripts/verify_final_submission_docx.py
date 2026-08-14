from __future__ import annotations

import hashlib
import zipfile
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
REPORT_SOURCE = Path(r"C:\Users\khiyw\Downloads\1780464077739_2026년도_인공지능_루키_본선_제안서.docx")
PLEDGE_SOURCE = Path(r"C:\Users\khiyw\Downloads\1780464099879_국내_AI_연계_기업_모델_활용_확약서.docx")
REPORT = ROOT / "output" / "docx" / "2026_AI_ROOKIE_본선제안서_Team_안전빵_작성본.docx"
PLEDGE = ROOT / "output" / "docx" / "국내_AI_연계_기업_모델_활용_확약서_Team_안전빵_작성본.docx"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def all_text(doc: Document) -> str:
    parts = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def main() -> None:
    assert sha256(REPORT_SOURCE) == "73F9042C5541FDBC647B04AD0C7CD305DD08780CBACBA8346ACE630D33AB2484"
    assert sha256(PLEDGE_SOURCE) == "176D6D1FB87BFFAB58F6D7EA75426C0A955697EA3B712A5F444FCF9696060F96"

    for path in (REPORT, PLEDGE):
        with zipfile.ZipFile(path) as archive:
            assert archive.testzip() is None, f"Corrupt DOCX package: {path}"

    report = Document(REPORT)
    report_text = all_text(report)
    assert len(report.sections) == 1
    assert len(report.tables) == 14
    assert len(report.inline_shapes) == 2
    for heading in (
        "1.1 도전과제명",
        "1.3 도전배경 및 목적",
        "1.4 기술 동향",
        "3.2.4 진척도",
        "4.1. 연구 지원 활용 항목",
        "6.4. 첨부 영상 설명",
        "7.4 국내 AI 모델 활용 시사점",
    ):
        assert heading in report_text, f"Missing heading: {heading}"
    for required in (
        "최신 승인 대표 시나리오(합성 기사 활용)",
        "활용 시간: 300시간",
        "모든 평가 내 숫자는 합성된 Live 자격평가입니다.",
        "실제 사고 및 현장 생산성 향상으로 해석하지 않습니다.",
        "영상 URL:",
        "해외(비생성)",
    ):
        assert required in report_text, f"Missing required text: {required}"
    for forbidden in (
        "합성 시연, 실제 사고확률 아님",
        "효과 주장 경계",
        "공개 Demo",
        "http://",
        "https://",
        "// 해당",
        "문상혁",
        "안재민",
    ):
        assert forbidden not in report_text, f"Forbidden text remains: {forbidden}"

    cover = report.tables[1]
    assert cover.cell(3, 1).text.rstrip().endswith("학년")
    education = report.tables[6]
    for row in education.rows[1:]:
        assert all(not row.cells[index].text.strip() for index in (2, 3, 4))
    video_paragraph = next(paragraph.text for paragraph in report.paragraphs if "영상 URL:" in paragraph.text)
    assert video_paragraph.rstrip().endswith("영상 URL:")

    pledge = Document(PLEDGE)
    pledge_text = all_text(pledge)
    assert len(pledge.sections) == 1
    assert len(pledge.tables) == 2
    model_table = pledge.tables[1].cell(0, 0).tables[0]
    assert len(model_table.rows) == 5
    assert [model_table.cell(row, 0).text.strip() for row in range(1, 5)] == [
        "업스테이지",
        "SKT",
        "SKT",
        "LG AI연구원",
    ]
    assert "2026년      월      일" in pledge_text
    assert "(서명 또는 인)" in pledge_text
    assert "김용우" in pledge_text and "안전빵" in pledge_text

    print("FINAL_DOCX_QA=PASS")
    print(f"REPORT_SHA256={sha256(REPORT)}")
    print(f"PLEDGE_SHA256={sha256(PLEDGE)}")


if __name__ == "__main__":
    main()
