from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


FONT = "맑은 고딕"
SOURCE_COLOR = RGBColor(89, 89, 89)


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def format_run(run, *, bold: bool = False) -> None:
    run.font.name = FONT
    run.font.size = Pt(7.2)
    run.font.bold = bold
    run.font.color.rgb = SOURCE_COLOR
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia"):
        rfonts.set(qn(f"w:{key}"), FONT)


def blank_before_heading(doc: Document, heading_prefix: str):
    for index, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip().startswith(heading_prefix):
            if index == 0 or doc.paragraphs[index - 1].text.strip():
                raise RuntimeError(f"No blank source slot before heading: {heading_prefix}")
            return doc.paragraphs[index - 1]
    raise RuntimeError(f"Heading not found: {heading_prefix}")


def set_source(paragraph, lines: list[tuple[str, str]]) -> None:
    clear_paragraph(paragraph)
    for line_index, (label, text) in enumerate(lines):
        if line_index:
            paragraph.add_run().add_break()
        label_run = paragraph.add_run(label)
        format_run(label_run, bold=True)
        text_run = paragraph.add_run(text)
        format_run(text_run)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(3)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    args.output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(args.input, args.output)
    doc = Document(args.output)

    source_13 = blank_before_heading(doc, "1.4 기술 동향")
    set_source(
        source_13,
        [
            (
                "출처: ",
                "장태원·김형렬·윤진하·강충원·이유민·민지희(2021), 「택배기사 적정 근로시간에 관한 연구」, 산업안전보건연구원.",
            )
        ],
    )

    source_14 = blank_before_heading(doc, "1.5 유사 기술 비교 및 차별점")
    set_source(
        source_14,
        [
            (
                "출처: ",
                "Useche et al.(2025), “The human cost of fast deliveries,” Journal of Transport & Health 44, 102133, DOI 10.1016/j.jth.2025.102133; Matre et al.(2021), “Safety incidents associated with extended working hours,” Scandinavian Journal of Work, Environment & Health 47(6), 415-424, DOI 10.5271/sjweh.3958.",
            ),
            (
                "추가 참고: ",
                "Fu & Ma(2022), Transportation Science 56(2), 404-435, DOI 10.1287/trsc.2021.1089; Rani·Pesole·González Vázquez(2024), EU·ILO, DOI 10.2760/712475; NIST(2023), AI RMF 1.0, DOI 10.6028/NIST.AI.100-1.",
            ),
        ],
    )

    doc.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
