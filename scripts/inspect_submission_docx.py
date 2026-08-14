from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph


def iter_blocks(parent):
    element = parent.element.body
    for child in element.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, parent)
        elif child.tag.endswith("}tbl"):
            yield Table(child, parent)


def inspect(path: Path) -> dict:
    doc = Document(path)
    blocks = []
    table_index = 0
    paragraph_index = 0
    for block in iter_blocks(doc):
        if isinstance(block, Paragraph):
            blocks.append(
                {
                    "type": "paragraph",
                    "index": paragraph_index,
                    "style": block.style.name if block.style else None,
                    "text": block.text,
                }
            )
            paragraph_index += 1
        else:
            rows = []
            for r, row in enumerate(block.rows):
                rows.append(
                    {
                        "row": r,
                        "cells": [
                            {
                                "col": c,
                                "text": cell.text,
                                "paragraphs": [p.text for p in cell.paragraphs],
                            }
                            for c, cell in enumerate(row.cells)
                        ],
                    }
                )
            blocks.append(
                {
                    "type": "table",
                    "index": table_index,
                    "rows": rows,
                }
            )
            table_index += 1

    def table_tree(table: Table) -> dict:
        return {
            "rows": [
                [
                    {
                        "text": cell.text,
                        "paragraphs": [p.text for p in cell.paragraphs],
                        "tables": [table_tree(nested) for nested in cell.tables],
                    }
                    for cell in row.cells
                ]
                for row in table.rows
            ]
        }

    return {
        "path": str(path),
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "inline_shape_count": len(doc.inline_shapes),
        "blocks": blocks,
        "table_trees": [table_tree(table) for table in doc.tables],
        "headers": [[p.text for p in s.header.paragraphs] for s in doc.sections],
        "footers": [[p.text for p in s.footer.paragraphs] for s in doc.sections],
    }


if __name__ == "__main__":
    source = Path(sys.argv[1])
    result = inspect(source)
    if len(sys.argv) > 2 and sys.argv[2] == "--summary":
        for block in result["blocks"]:
            if block["type"] == "paragraph" and block["text"]:
                print(f"P{block['index']:03}: {block['text']!r}")
            elif block["type"] == "table":
                print(f"TABLE {block['index']}")
                for row in block["rows"]:
                    print("  " + " || ".join(cell["text"].replace("\n", " / ") for cell in row["cells"]))
    else:
        print(json.dumps(result, ensure_ascii=False, indent=2))
